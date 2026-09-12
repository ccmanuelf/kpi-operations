"""The generators must actually render the configured target, not a literal.

`test_report_targets.py` gates the resolver. These gate the wiring: that a
target an administrator configured reaches the cell a reader sees, in both
formats, and that the two formats agree about the same client.

Each of the five PDF summary rows had its own literal, so each is asserted
separately -- a single test on efficiency would have passed while FPY still
printed 99.
"""

from datetime import date
from decimal import Decimal

import pytest

from backend.db.factories import TestDataFactory
from backend.orm.kpi_threshold import KPIThreshold
from backend.reports.excel_generator import ExcelReportGenerator
from backend.reports.pdf_generator import PDFReportGenerator

DAY = date(2026, 6, 11)


def _threshold(db, *, client_id, kpi_key, target, warning=None, critical=None, unit="%", higher="Y"):
    db.add(
        KPIThreshold(
            threshold_id=f"THR-{client_id or 'GLOBAL'}-{kpi_key}",
            client_id=client_id,
            kpi_key=kpi_key,
            target_value=target,
            warning_threshold=warning,
            critical_threshold=critical,
            unit=unit,
            higher_is_better=higher,
        )
    )
    db.flush()


@pytest.fixture
def tenant(transactional_db):
    """One client with a production, quality and attendance row for DAY."""
    db = transactional_db
    client = TestDataFactory.create_client(db)
    user = TestDataFactory.create_user(db, client_id=client.client_id)
    shift = TestDataFactory.create_shift(db, client_id=client.client_id)
    product = TestDataFactory.create_product(db, client_id=client.client_id)
    employee = TestDataFactory.create_employee(db, client_id=client.client_id)
    work_order = TestDataFactory.create_work_order(db, client_id=client.client_id, product_id=product.product_id)
    # efficiency_percentage / performance_percentage are set explicitly because
    # nothing in the application writes them, and the generators now OMIT a row
    # whose measurement was never recorded (backend/reports/measurements.py).
    # These tests are about which TARGET reaches the cell, so the rows have to
    # render -- an unmeasured fixture would make every assertion below vacuous.
    TestDataFactory.create_production_entry(
        db,
        client_id=client.client_id,
        product_id=product.product_id,
        shift_id=shift.shift_id,
        entered_by=user.user_id,
        production_date=DAY,
        efficiency_percentage=Decimal("80.0"),
        performance_percentage=Decimal("90.0"),
    )
    TestDataFactory.create_quality_entry(
        db,
        work_order_id=work_order.work_order_id,
        client_id=client.client_id,
        inspector_id=user.user_id,
        inspection_date=DAY,
    )
    TestDataFactory.create_attendance_entry(
        db,
        employee_id=employee.employee_id,
        client_id=client.client_id,
        shift_id=shift.shift_id,
        shift_date=DAY,
        scheduled_hours=Decimal("8.0"),
        absence_hours=Decimal("1.0"),
        is_absent=True,
    )
    db.commit()
    return db, client


def _pdf_row(db, client_id, name):
    rows = PDFReportGenerator(db)._fetch_kpi_summary(client_id, DAY, DAY)
    return next((r for r in rows if r["name"] == name), None)


def _excel_row(db, client_id, name):
    rows = ExcelReportGenerator(db)._fetch_kpi_summary_data(client_id, DAY, DAY)
    return next((r for r in rows if r["name"] == name), None)


# (PDF row name, excel row name, kpi_key, a configured target, its old literal)
ROWS = [
    ("Efficiency", "Efficiency", "efficiency", 42.0, 85),
    ("Performance", "Performance", "performance", 43.0, 85),
    ("First Pass Yield", "FPY", "fpy", 44.0, 99),
    ("PPM", "PPM", "ppm", 45.0, 1000),
    ("Absenteeism", "Absenteeism", "absenteeism", 46.0, 5),
]


class TestTheConfiguredTargetReachesTheCell:
    @pytest.mark.parametrize("pdf_name,excel_name,kpi_key,configured,old_literal", ROWS)
    def test_pdf(self, tenant, pdf_name, excel_name, kpi_key, configured, old_literal):
        db, client = tenant
        _threshold(db, client_id=client.client_id, kpi_key=kpi_key, target=configured)
        db.commit()

        row = _pdf_row(db, client.client_id, pdf_name)
        assert row is not None, f"{pdf_name} row absent"
        assert row["target"] == configured, f"{pdf_name} ignored configuration (literal was {old_literal})"

    @pytest.mark.parametrize("pdf_name,excel_name,kpi_key,configured,old_literal", ROWS)
    def test_excel(self, tenant, pdf_name, excel_name, kpi_key, configured, old_literal):
        db, client = tenant
        _threshold(db, client_id=client.client_id, kpi_key=kpi_key, target=configured)
        db.commit()

        row = _excel_row(db, client.client_id, excel_name)
        assert row is not None, f"{excel_name} row absent"
        assert row["target"] == configured, f"{excel_name} ignored configuration (literal was {old_literal})"


class TestTheGlobalRowIsTheFallback:
    def test_a_client_without_its_own_row_reads_the_global_one(self, tenant):
        db, client = tenant
        _threshold(db, client_id=None, kpi_key="efficiency", target=51.0)
        db.commit()

        assert _pdf_row(db, client.client_id, "Efficiency")["target"] == 51.0

    def test_the_clients_own_row_wins(self, tenant):
        db, client = tenant
        _threshold(db, client_id=None, kpi_key="efficiency", target=51.0)
        _threshold(db, client_id=client.client_id, kpi_key="efficiency", target=52.0)
        db.commit()

        assert _pdf_row(db, client.client_id, "Efficiency")["target"] == 52.0


class TestBothFormatsAgree:
    def test_pdf_and_excel_report_the_same_target_and_status(self, tenant):
        # The divergence #300 fixed was PDF and Excel disagreeing about one
        # client's absenteeism. Targets must not reintroduce it.
        db, client = tenant
        _threshold(db, client_id=client.client_id, kpi_key="efficiency", target=99.0, warning=50.0, critical=25.0)
        db.commit()

        pdf = _pdf_row(db, client.client_id, "Efficiency")
        excel = _excel_row(db, client.client_id, "Efficiency")
        assert (pdf["target"], pdf["status"]) == (excel["target"], excel["status"])


def _delete_global(db, kpi_key):
    """Remove the global default migration 0009 supplies.

    Every metric the report shows now has one, so an unconfigured metric is
    reachable only if an administrator deletes the global row -- which the
    threshold editor permits. The report still has to render.
    """
    db.query(KPIThreshold).filter(KPIThreshold.client_id.is_(None), KPIThreshold.kpi_key == kpi_key).delete()
    db.commit()


class TestAnUnconfiguredMetric:
    def test_the_pdf_reports_no_target_rather_than_a_literal(self, tenant):
        db, client = tenant
        _delete_global(db, "absenteeism")

        row = _pdf_row(db, client.client_id, "Absenteeism")
        assert row["target"] is None
        assert row["status"] == "No Target"

    def test_the_pdf_renders_it_as_an_em_dash(self, tenant):
        # The status is only honest if the Target cell stops showing a number.
        db, client = tenant
        _delete_global(db, "absenteeism")
        elements = PDFReportGenerator(db)._build_executive_summary(client.client_id, DAY, DAY)

        tables = [e for e in elements if hasattr(e, "_cellvalues")]
        assert tables, "no summary table rendered"
        cells = [str(c) for row in tables[0]._cellvalues for c in row]
        assert "—" in cells, f"expected an em dash in the Target column: {cells}"

    def test_excel_leaves_the_target_cell_blank(self, tenant):
        db, client = tenant
        _delete_global(db, "absenteeism")

        row = _excel_row(db, client.client_id, "Absenteeism")
        assert row["target"] is None
        assert row["status"] == "No Target"

    def test_excel_leaves_the_variance_blank_too(self, tenant):
        # Excel reads an empty Target cell as 0, so `=B-C` against a blank target
        # renders the measured value as its own variance -- a wrong number rather
        # than an absent one.
        db, client = tenant
        _delete_global(db, "absenteeism")
        from openpyxl import load_workbook

        buf = ExcelReportGenerator(db).generate_report(client_id=client.client_id, start_date=DAY, end_date=DAY)
        ws = load_workbook(buf)["Executive Summary"]

        for row in range(8, ws.max_row + 1):
            if ws[f"A{row}"].value == "Absenteeism":
                assert ws[f"C{row}"].value is None, "target should be blank"
                assert ws[f"D{row}"].value is None, f"variance holds {ws[f'D{row}'].value!r}"
                break
        else:
            raise AssertionError("no Absenteeism row in the summary sheet")

    def test_excel_does_not_paint_it_as_an_alarm(self, tenant):
        # The status-colour rule falls through to red for anything it does not
        # recognise, so an unconfigured target would have rendered as critical.
        from backend.reports.targets import UNEVALUATED_STATUSES

        db, client = tenant
        _delete_global(db, "absenteeism")
        row = _excel_row(db, client.client_id, "Absenteeism")
        assert row["status"] in UNEVALUATED_STATUSES


class TestTheTrendColumnIsGone:
    def test_no_excel_row_carries_a_trend(self, tenant):
        # It restated a threshold against a literal 85, six values were a
        # hardcoded arrow, and the PPM arrow rose as the metric worsened.
        db, client = tenant
        rows = _excel_row(db, client.client_id, "Efficiency")
        assert "trend" not in rows

    def test_the_summary_sheet_header_stops_at_status(self, tenant):
        db, client = tenant
        from openpyxl import load_workbook

        buf = ExcelReportGenerator(db).generate_report(client_id=client.client_id, start_date=DAY, end_date=DAY)
        wb = load_workbook(buf)
        assert "Executive Summary" in wb.sheetnames, wb.sheetnames
        ws = wb["Executive Summary"]
        assert ws["E7"].value == "Status"
        assert ws["F7"].value is None, f"F7 still holds {ws['F7'].value!r}"

    def test_the_summary_sheet_does_not_still_SHOW_a_sixth_column(self, tenant):
        """Removing the data is not removing the column.

        Four places kept drawing column F after the Trend values were gone: the
        alternating-row fill striped it, `_apply_table_borders` ran the table out
        to F, a width was reserved for it, and the title banner was merged across
        A1:F1. The sheet would have rendered a bordered, striped, empty sixth
        column under a banner sized for it -- which reads as a field that failed
        to populate rather than one that was removed.

        Checks MEMBERSHIP of column_dimensions rather than reading
        `ws.column_dimensions["F"].width`: openpyxl creates a default dimension
        on access, so reading it fabricates the very thing being asserted about.
        """
        db, client = tenant
        from openpyxl import load_workbook

        buf = ExcelReportGenerator(db).generate_report(client_id=client.client_id, start_date=DAY, end_date=DAY)
        ws = load_workbook(buf)["Executive Summary"]

        assert "F" not in ws.column_dimensions, "column F still has a reserved width"
        assert (
            str(ws.merged_cells.ranges) == "{<MergedCellRange A1:E1>}"
        ), f"the title banner still spans the removed column: {ws.merged_cells.ranges}"
        for row in range(7, ws.max_row + 1):
            cell = ws[f"F{row}"]
            assert cell.value is None, f"F{row} holds {cell.value!r}"
            assert cell.border.left.style is None, f"F{row} is still bordered"
            assert cell.fill.fill_type in (None, "none"), f"F{row} is still filled"

    def test_the_pdf_detail_block_carries_no_trend_key(self, tenant):
        db, client = tenant
        details = PDFReportGenerator(db)._fetch_kpi_details("efficiency", client.client_id, DAY, DAY)
        assert "Trend" not in details, details


class TestTheStatusEscalatesOnConfiguredBands:
    """The seeded row is 1 absent hour in 8 scheduled, so the value is 12.5%.

    Absenteeism is lower-is-better, so the bands ascend past the target. The
    numbers below isolate one verdict each: `check_threshold_breach`'s one
    unconditional rule fires once the value exceeds five times target, which
    would otherwise make every case Urgent.
    """

    def _configure(self, db, client, *, target, warning, critical):
        _threshold(
            db,
            client_id=client.client_id,
            kpi_key="absenteeism",
            target=target,
            warning=warning,
            critical=critical,
            higher="N",
        )
        db.commit()

    def test_below_target_but_inside_both_bands_is_at_risk(self, tenant):
        db, client = tenant
        self._configure(db, client, target=10.0, warning=20.0, critical=30.0)
        assert _pdf_row(db, client.client_id, "Absenteeism")["status"] == "At Risk"

    def test_breaching_the_warning_band_reads_warning(self, tenant):
        db, client = tenant
        self._configure(db, client, target=10.0, warning=12.0, critical=30.0)
        assert _pdf_row(db, client.client_id, "Absenteeism")["status"] == "Warning"

    def test_breaching_the_critical_band_reads_critical(self, tenant):
        db, client = tenant
        self._configure(db, client, target=10.0, warning=11.0, critical=12.0)
        assert _pdf_row(db, client.client_id, "Absenteeism")["status"] == "Critical"

    def test_far_past_target_reads_urgent(self, tenant):
        db, client = tenant
        self._configure(db, client, target=2.0, warning=3.0, critical=4.0)
        assert _pdf_row(db, client.client_id, "Absenteeism")["status"] == "Urgent"

    def test_meeting_target_reads_on_target(self, tenant):
        db, client = tenant
        self._configure(db, client, target=20.0, warning=30.0, critical=40.0)
        assert _pdf_row(db, client.client_id, "Absenteeism")["status"] == "On Target"


class TestTheDetailBlocksTargetsToo:
    """The summary was not the only place with literals.

    One detail block serves efficiency, performance AND availability off a
    single hardcoded "85%", so two of the three showed a target that was not
    theirs -- performance's global default is 95 and availability's is 90.
    """

    @pytest.mark.parametrize(
        "kpi_key,configured", [("efficiency", 61.0), ("performance", 62.0), ("availability", 63.0)]
    )
    def test_each_of_the_three_shared_metrics_shows_its_own_target(self, tenant, kpi_key, configured):
        db, client = tenant
        _threshold(db, client_id=client.client_id, kpi_key=kpi_key, target=configured)
        db.commit()

        details = PDFReportGenerator(db)._fetch_kpi_details(kpi_key, client.client_id, DAY, DAY)
        assert details.get("Target") == f"{configured:g}%", details

    def test_the_three_do_not_share_one_number(self, tenant):
        # The shape of the original bug: distinct configuration, one rendered
        # value. Asserting they DIFFER catches a reversion that a single
        # per-metric assertion would not.
        db, client = tenant
        for key, value in (("efficiency", 61.0), ("performance", 62.0), ("availability", 63.0)):
            _threshold(db, client_id=client.client_id, kpi_key=key, target=value)
        db.commit()

        gen = PDFReportGenerator(db)
        shown = {
            k: gen._fetch_kpi_details(k, client.client_id, DAY, DAY).get("Target")
            for k in ("efficiency", "performance", "availability")
        }
        assert len(set(shown.values())) == 3, shown

    def test_the_variance_is_measured_against_the_configured_target(self, tenant):
        db, client = tenant
        _threshold(db, client_id=client.client_id, kpi_key="efficiency", target=10.0)
        db.commit()

        details = PDFReportGenerator(db)._fetch_kpi_details("efficiency", client.client_id, DAY, DAY)
        current = float(details["Current Value"].rstrip("%"))
        assert details["Variance"] == f"{current - 10.0:+.1f}%", details

    @pytest.mark.parametrize("kpi_key", ["fpy", "ppm", "absenteeism"])
    def test_the_quality_and_attendance_details_read_configuration(self, tenant, kpi_key):
        db, client = tenant
        _threshold(db, client_id=client.client_id, kpi_key=kpi_key, target=64.0)
        db.commit()

        details = PDFReportGenerator(db)._fetch_kpi_details(kpi_key, client.client_id, DAY, DAY)
        # PPM renders bare, the percentages carry a suffix -- as they always have.
        assert details.get("Target") in ("64%", "64"), details

    def test_an_unconfigured_detail_target_is_an_em_dash(self, tenant):
        db, client = tenant
        _delete_global(db, "absenteeism")

        details = PDFReportGenerator(db)._fetch_kpi_details("absenteeism", client.client_id, DAY, DAY)
        assert details.get("Target") == "—", details
